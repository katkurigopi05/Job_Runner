"""The owner's company workbook: every company, and the ones without a readable board.

Asked for on 2026-10-19. The owner's sheet lists companies; the crawler reads
jobs from the supported boards in `BOARDS`. The second sheet lists companies
whose boards have not been verified yet, with what is known about each:
the careers page discovery found on its website, where the deep dive found
applicants are sent, and what to do. It is rebuilt from the database
(`make company-sheet`), so it is the current state of the same list the
workers act on — discovery and the deep dive take exactly the companies with
no board — rather than a copy that drifts.

Read-only: building it changes no company, note or task.
"""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

BOARDS = {
    "greenhouse": "Greenhouse",
    "ashby": "Ashby",
    "lever": "Lever",
    "workable": "Workable",
    "workday": "Workday",
}
MISSING_SHEET = "Not on a job board we read"

#: How each company came to be on a board, by the evidence's `method`.
FOUND_BY = {
    "discovery_from_directory": "board directory",
    "discovery_from_url": "its website",
    "discovery_from_name": "its name",
    "registry_sync": "registry",
    "deep_dive": "deep dive",
}

#: What discovery saw, in the owner's words (`find_boards.SiteReading.category`).
WHAT_WE_FOUND = {
    "other_job_system": "Uses another job system",
    "careers_page_unreadable": "Careers page found, no job board on it",
    "no_careers_page": "No careers page found on its website",
    "unreachable": "Website did not answer",
    "blocked": "Website blocks automated reading (robots.txt)",
    "multi_company_board": "Careers page lists many companies' jobs",
    "structured_jobs": "Careers page lists jobs as structured data",
    "board_empty": "Job board found, but no open jobs",
    "no_website": "No website on record",
}

DEEP_DIVE_STATE = {"noted": "Not started", "queued": "Waiting", "needs_owner": "Done"}

ALL_COLUMNS = (
    "Company Name",
    "Website URL",
    "Jobs/Careers URL",
    "On a job board?",
    "Job board",
    "Board URL",
    "Open jobs",
    "Found by",
)
MISSING_COLUMNS = (
    "Company Name",
    "Website URL",
    "Careers page",
    "Where to apply",
    "What we found",
    "Deep dive",
    "Deep dive found",
    "What to do",
    "Details",
    "Original careers link (from CSV)",
)

NAME, WEBSITE, CAREERS = "Company Name", "Website URL", "Jobs/Careers URL"


@dataclass(frozen=True)
class Company:
    """One company row and its discovery note, as the sheet needs them."""

    name: str
    status: str
    ats_type: str | None
    slug: str | None
    board_url: str | None
    evidence: dict[str, Any] = field(default_factory=dict)
    reason: str | None = None
    holder: str | None = None
    open_jobs: int = 0
    issue_status: str | None = None
    findings: dict[str, Any] = field(default_factory=dict)
    diagnosis: str | None = None
    fix: str | None = None


def _board(ats: str | None) -> str:
    return BOARDS.get(ats or "", ats or "")


def _without_guessing(reason: str | None) -> str:
    return re.sub(r";\s*name-guessing is off.*$", "", reason or "").strip()


def all_row(row: Mapping[str, str], company: Company | None) -> list[Any]:
    """The first sheet's row: the CSV's columns, then where its jobs come from."""
    base = [row.get(NAME, ""), row.get(WEBSITE, ""), row.get(CAREERS, "")]
    if company is None:
        return [*base, "Not imported", "", "", "", ""]
    if company.status == "verified":
        found_by = FOUND_BY.get(company.evidence.get("method", ""), "")
        return [
            *base,
            "Yes",
            _board(company.ats_type),
            company.board_url or "",
            company.open_jobs,
            found_by,
        ]
    if company.evidence.get("method") == "duplicate_board":
        ats, slug = company.evidence.get("ats", ""), company.evidence.get("slug", "")
        holder = company.holder or "another company"
        return [*base, f"Same board as {holder}", _board(ats), f"{ats}/{slug}", "", ""]
    return [*base, "No", "", "", "", ""]


def careers_page(company: Company) -> str:
    """The best careers page known: discovery's, else one the deep dive read."""
    findings = company.findings
    return str(
        company.evidence.get("careers_page")
        or (findings.get("careers_page") or {}).get("url")
        or (findings.get("further") or {}).get("url")
        or ""
    )


def where_to_apply(company: Company) -> str:
    """Where applications actually go, if anything has said."""
    findings = company.findings
    for candidate in (
        company.evidence.get("job_system_url"),
        *(findings.get("other_sites") or [])[:1],
        *(findings.get("own_postings") or [])[:1],
    ):
        if candidate:
            return str(candidate)
    emails = findings.get("emails") or []
    return f"mailto:{emails[0]}" if emails else ""


def missing_row(row: Mapping[str, str], company: Company | None) -> list[str] | None:
    """The second sheet's row, or None for a company whose jobs are on a board."""
    if company is not None and (
        company.status == "verified" or company.evidence.get("method") == "duplicate_board"
    ):
        return None
    if company is None:
        what, deep, page, apply, details = "Not imported yet", "", "", "", ""
        diagnosis = fix = ""
    else:
        category = company.evidence.get("category", "")
        what = WHAT_WE_FOUND.get(category, "Not yet checked with the latest search")
        system = company.evidence.get("job_system")
        if str(system or "").strip().casefold() == "workday":
            what = "Workday: readable now, waiting for re-discovery"
        elif system:
            what = f"{what} ({system})"
        deep = DEEP_DIVE_STATE.get(company.issue_status or "noted", "Not started")
        page, apply = careers_page(company), where_to_apply(company)
        details = _without_guessing(company.reason)
        diagnosis, fix = company.diagnosis or "", company.fix or ""
    return [
        row.get(NAME, ""),
        row.get(WEBSITE, ""),
        page,
        apply,
        what,
        deep,
        diagnosis,
        fix,
        details,
        row.get(CAREERS, ""),
    ]


async def load_companies(session: AsyncSession, source_file: str) -> dict[int, Company]:
    """Every company imported from `source_file`, by its row in that file."""
    rows = (
        await session.execute(
            text("""
            SELECT c.source_row, c.name, c.source_status, c.ats_type, c.slug, c.careers_url,
                   c.source_evidence AS evidence, c.discovery_failure, h.name AS holder,
                   (SELECT count(*) FROM postings p
                     WHERE p.company_id = c.id AND p.closed_at IS NULL) AS open_jobs,
                   i.status AS issue_status, i.findings, i.diagnosis, i.fix
            FROM companies c
            LEFT JOIN companies h ON h.id::text = c.source_evidence->>'duplicate_of'
            LEFT JOIN LATERAL (
                SELECT status, findings, diagnosis, fix FROM company_issues
                WHERE company_id = c.id AND stage = 'discovery'
                ORDER BY (status <> 'fixed') DESC, last_noted_at DESC
                LIMIT 1
            ) i ON true
            WHERE c.source_file = :source_file AND c.source_row IS NOT NULL
            """),
            {"source_file": source_file},
        )
    ).mappings()
    return {
        int(r["source_row"]): Company(
            name=r["name"],
            status=r["source_status"],
            ats_type=r["ats_type"],
            slug=r["slug"],
            board_url=r["careers_url"],
            evidence=r["evidence"] or {},
            reason=r["discovery_failure"],
            holder=r["holder"],
            open_jobs=int(r["open_jobs"] or 0),
            issue_status=r["issue_status"],
            findings=r["findings"] or {},
            diagnosis=r["diagnosis"],
            fix=r["fix"],
        )
        for r in rows
    }


def write_workbook(
    path: Path,
    *,
    all_rows: Sequence[Sequence[Any]],
    missing_rows: Sequence[Sequence[Any]],
    source_name: str,
    built_at: datetime | None = None,
) -> None:
    """Write the two sheets and a short note. Arial, header frozen, filterable."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    first = book.active
    assert first is not None
    first.title = "All companies"
    second = book.create_sheet(MISSING_SHEET)
    for sheet, columns, rows, widths in (
        (first, ALL_COLUMNS, all_rows, (34, 34, 46, 26, 13, 52, 10, 16)),
        (second, MISSING_COLUMNS, missing_rows, (34, 34, 46, 46, 40, 12, 60, 40, 60, 40)),
    ):
        sheet.append(list(columns))
        for row in rows:
            sheet.append(list(row))
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
            cell.alignment = Alignment(vertical="center")
        body = Font(name="Arial", size=10)
        for line in sheet.iter_rows(min_row=2):
            for cell in line:
                cell.font = body
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    note = book.create_sheet("About")
    when = (built_at or datetime.now()).strftime("%Y-%m-%d %H:%M")
    for line in (
        "About this workbook",
        f"Rebuilt {when} from {source_name} and the Jobrunner database (make company-sheet).",
        "All companies: every row of the CSV, unchanged, then whether its jobs are on a board "
        "we read, how many are open, and how the board was found.",
        f"Job boards we read: {', '.join(BOARDS.values())}.",
        f"{MISSING_SHEET}: the companies whose boards have not been verified yet — the same list "
        "discovery and the deep dive work on — with the careers page found on each website, where "
        "the deep dive found applicants are sent, and what to do.",
        "'Same board as X': this company's board is already read under another company row.",
        "Rebuild it after a crawl or a deep dive; the CSV itself is never changed.",
    ):
        note.append([line])
    note["A1"].font = Font(name="Arial", bold=True, size=12)
    for line in note.iter_rows(min_row=2):
        for cell in line:
            cell.font = Font(name="Arial", size=10)
    note.column_dimensions["A"].width = 120
    book.save(path)


__all__ = [
    "ALL_COLUMNS",
    "MISSING_COLUMNS",
    "Company",
    "all_row",
    "careers_page",
    "load_companies",
    "missing_row",
    "where_to_apply",
    "write_workbook",
]
