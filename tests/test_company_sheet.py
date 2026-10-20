"""`make company-sheet`: the owner's companies, and the ones not on the 4 job boards.

Asked for on 2026-10-19: a second sheet beside the CSV listing the companies
whose jobs are not on Greenhouse, Ashby, Lever or Workable, with the careers
pages discovery and the deep dive found, kept current so it is the list the
owner works from.
"""

from __future__ import annotations

import csv
import sys

from openpyxl import load_workbook

from packages.core.enums import SourceStatus
from packages.core.models import Company as CompanyRow
from packages.core.models_issues import CompanyIssue
from packages.crawler import company_sheet as sheet
from packages.crawler.company_sheet import Company

ROW = {"Company Name": "Acme", "Website URL": "https://acme.com", "Jobs/Careers URL": "https://g/q"}


def _company(**values) -> Company:
    defaults = {
        "name": "Acme",
        "status": "failed",
        "ats_type": None,
        "slug": None,
        "board_url": None,
        "evidence": {},
        "reason": None,
        "holder": None,
        "open_jobs": 0,
        "issue_status": None,
        "findings": {},
        "diagnosis": None,
        "fix": None,
    }
    return Company(**{**defaults, **values})


# --- the first sheet: every company ---------------------------------------------------


def test_a_company_on_a_job_board_says_which_and_is_not_on_the_second_sheet() -> None:
    acme = _company(
        status="verified",
        ats_type="greenhouse",
        slug="acme",
        board_url="https://boards.greenhouse.io/acme",
        evidence={"method": "discovery_from_url"},
        open_jobs=12,
    )

    assert sheet.all_row(ROW, acme)[3:] == [
        "Yes",
        "Greenhouse",
        "https://boards.greenhouse.io/acme",
        12,
        "its website",
    ]
    assert sheet.missing_row(ROW, acme) is None


def test_a_board_found_by_the_deep_dive_says_so() -> None:
    acme = _company(status="verified", ats_type="ashby", evidence={"method": "deep_dive"})

    assert sheet.all_row(ROW, acme)[-1] == "deep dive"


def test_a_company_whose_board_another_row_holds_names_it() -> None:
    zeit = _company(
        evidence={"method": "duplicate_board", "ats": "greenhouse", "slug": "vercel"},
        holder="Vercel",
    )

    assert sheet.all_row(ROW, zeit)[3:6] == [
        "Same board as Vercel",
        "Greenhouse",
        "greenhouse/vercel",
    ]
    assert sheet.missing_row(ROW, zeit) is None


def test_a_row_never_imported_is_listed_as_such() -> None:
    assert sheet.all_row(ROW, None)[3] == "Not imported"
    assert sheet.missing_row(ROW, None)[sheet.MISSING_COLUMNS.index("What we found")] == (
        "Not imported yet"
    )


# --- the second sheet: the companies not on the 4 job boards ---------------------------


def _cell(row: list, column: str):
    return row[sheet.MISSING_COLUMNS.index(column)]


def test_what_discovery_found_is_carried() -> None:
    acme = _company(
        evidence={
            "category": "other_job_system",
            "careers_page": "https://acme.com/careers",
            "job_system": "Workday",
            "job_system_url": "https://acme.wd5.myworkdayjobs.com/External",
        },
        reason="careers page uses Workday; name-guessing is off for a company with a website",
        issue_status="queued",
    )

    row = sheet.missing_row(ROW, acme)

    assert _cell(row, "Careers page") == "https://acme.com/careers"
    assert _cell(row, "Where to apply") == "https://acme.wd5.myworkdayjobs.com/External"
    assert _cell(row, "What we found") == "Uses another job system (Workday)"
    assert _cell(row, "Deep dive") == "Waiting"
    assert "name-guessing" not in _cell(row, "Details")
    assert _cell(row, "Original careers link (from CSV)") == "https://g/q"


def test_what_the_deep_dive_found_fills_the_gaps() -> None:
    acme = _company(
        evidence={"category": "no_careers_page"},
        issue_status="needs_owner",
        findings={
            "kind": "sends_elsewhere",
            "further": {"url": "https://acme.com/jobs-at-acme"},
            "other_sites": ["https://careers.risepeople.com/acme/1"],
        },
        diagnosis="its careers page sends applicants to careers.risepeople.com (...)",
        fix="look at careers.risepeople.com",
    )

    row = sheet.missing_row(ROW, acme)

    assert _cell(row, "Careers page") == "https://acme.com/jobs-at-acme"
    assert _cell(row, "Where to apply") == "https://careers.risepeople.com/acme/1"
    assert _cell(row, "Deep dive") == "Done"
    assert _cell(row, "Deep dive found").startswith("its careers page sends applicants")
    assert _cell(row, "What to do") == "look at careers.risepeople.com"


def test_an_email_application_is_where_to_apply() -> None:
    acme = _company(issue_status="needs_owner", findings={"emails": ["jobs@acme.com"]})

    assert _cell(sheet.missing_row(ROW, acme), "Where to apply") == "mailto:jobs@acme.com"


def test_a_company_not_looked_into_says_so() -> None:
    assert _cell(sheet.missing_row(ROW, _company()), "Deep dive") == "Not started"


# --- the workbook ---------------------------------------------------------------------


def test_the_workbook_has_both_sheets_and_a_note(tmp_path) -> None:
    out = tmp_path / "companies.xlsx"
    on_board = _company(status="verified", ats_type="lever", open_jobs=3)

    sheet.write_workbook(
        out,
        all_rows=[sheet.all_row(ROW, on_board), sheet.all_row(ROW, _company())],
        missing_rows=[sheet.missing_row(ROW, _company())],
        source_name="companies.csv",
    )

    book = load_workbook(out)
    assert book.sheetnames == ["All companies", "Not on the 4 job boards", "About"]
    first, second = book["All companies"], book["Not on the 4 job boards"]
    assert [c.value for c in first[1]] == list(sheet.ALL_COLUMNS)
    assert [c.value for c in second[1]] == list(sheet.MISSING_COLUMNS)
    assert (first.max_row, second.max_row) == (3, 2)
    assert first.freeze_panes == "A2" and second.auto_filter.ref


async def test_the_command_reads_the_database_and_writes_beside_the_csv(
    committing_sessionmaker, worker_session, monkeypatch, tmp_path, capsys
) -> None:
    from scripts import company_sheet

    source = tmp_path / "companies.csv"
    with source.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Company Name", "Website URL", "Jobs/Careers URL"])
        writer.writerow(["Onboard", "https://onboard.com", ""])
        writer.writerow(["Offboard", "https://offboard.com", ""])
    on = CompanyRow(
        name="Onboard",
        source_file="companies.csv",
        source_row=1,
        source_status=SourceStatus.VERIFIED.value,
        ats_type="ashby",
        slug="onboard",
    )
    off = CompanyRow(
        name="Offboard",
        source_file="companies.csv",
        source_row=2,
        source_status=SourceStatus.FAILED.value,
        source_evidence={"category": "unreachable"},
    )
    worker_session.add_all([on, off])
    await worker_session.flush()
    worker_session.add(
        CompanyIssue(
            company_id=off.id,
            stage="discovery",
            category="unreachable",
            reason="r",
            status="needs_owner",
            findings={"kind": "domain_gone"},
            diagnosis="its domain no longer exists",
        )
    )
    await worker_session.commit()
    monkeypatch.setattr(company_sheet, "get_sessionmaker", lambda: committing_sessionmaker)
    monkeypatch.setattr(sys, "argv", ["company_sheet", str(source)])

    assert await company_sheet.main() == 0

    book = load_workbook(tmp_path / "companies.xlsx")
    rows = list(book["Not on the 4 job boards"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["Offboard"]
    assert rows[0][sheet.MISSING_COLUMNS.index("Deep dive found")] == "its domain no longer exists"
    assert "1 on a job board, 1 not" in capsys.readouterr().out
